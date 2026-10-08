__all__ = [
    "compute_file_hash",
    "embed_chunks_with_retry",
    "expand_user_roles",
    "get_db_connection",
    "get_documents",
    "get_embedding_model",
    "get_llm",
    "get_ollama_host",
    "get_retry_attempts",
    "get_timeout_seconds",
    "load_config",
    "log_error",
    "remove_comments_and_docstrings",
    "rerank_candidates",
    "safe_calculate",
    "split_chunks",
    "truncate_and_normalize",
]

import ast
import hashlib
import logging
import operator
import os
import re
from pathlib import Path
from typing import Union

import numpy as np
import openai
import psycopg2
from dotenv import load_dotenv
from langchain_community.document_loaders import Docx2txtLoader, PyPDFLoader, UnstructuredMarkdownLoader
from langchain_core.documents import Document
from langchain_openai import ChatOpenAI, OpenAIEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter
from tenacity import retry, retry_if_exception, stop_after_attempt, wait_exponential

ROLE_HIERARCHY = {
    "admin": ["admin", "executive", "finance_analyst", "engineering"],
    "finance_lead": ["finance_lead", "finance_analyst"],
    "engineering": ["engineering"]
}

_SAFE_OPERATORS = {
    ast.Add: operator.add,
    ast.Sub: operator.sub,
    ast.Mult: operator.mul,
    ast.Div: operator.truediv,
    ast.Pow: operator.pow,
    ast.USub: operator.neg,
}

_DOCSTRING_PATTERN = re.compile(r'"""[\s\S]*?"""')

def expand_user_roles(raw_roles: list[str]) -> list[str]:
    expanded = set()
    for role in raw_roles:
        expanded.update(ROLE_HIERARCHY.get(role, [role]))
    return list(expanded)

def compute_file_hash(file_path: Union[str, Path]) -> str:
    """
    Compute SHA-256 hash of a file in 64KB chunks for memory efficiency.
    Used for cryptographic de-duplication in the ingestion pipeline.
    """
    hasher = hashlib.sha256()
    with open(file_path, "rb") as f:
        while chunk := f.read(65536):
            hasher.update(chunk)
    return hasher.hexdigest()


def truncate_and_normalize(
    embedding: Union[list[float], "np.ndarray"],
    target_dim: int = 1536
) -> list[float]:
    vec = np.asarray(embedding[:target_dim], dtype=np.float32)
    norm = np.linalg.norm(vec)
    if norm == 0.0:
        return vec.tolist()
    normalized_vec = vec / norm
    return normalized_vec.tolist()


def load_config(env_path: str | None = None) -> tuple[str, str]:
    if env_path is None:
        env_path = str(Path(__file__).parent / ".env")
    load_dotenv(env_path)
    openai_api_key = os.getenv("OPENAI_API_KEY", "")
    database_url = os.getenv("DATABASE_URL", "")
    if not openai_api_key:
        raise ValueError("OPENAI_API_KEY missing from .env.")
    if not database_url:
        raise ValueError("DATABASE_URL missing from .env.")
    return openai_api_key, database_url


def get_documents(file_paths: list[str]) -> list[Document]:
    print(f"📄 Loading {len(file_paths)} document(s)...")
    documents = []
    for file_path in file_paths:
        try:
            suffix = Path(file_path).suffix.lower()
            if suffix == ".pdf":
                loader = PyPDFLoader(file_path)
                documents.extend(loader.load())
            elif suffix == ".md":
                loader = UnstructuredMarkdownLoader(file_path)
                documents.extend(loader.load())
            elif suffix == ".docx":
                try:
                    loader = Docx2txtLoader(file_path)
                    documents.extend(loader.load())
                except (ImportError, ModuleNotFoundError):
                    import docx

                    doc = docx.Document(file_path)
                    content_parts = [p.text for p in doc.paragraphs if p.text.strip()]
                    for table in doc.tables:
                        for row in table.rows:
                            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                            if row_text:
                                content_parts.append(row_text)
                    text = "\n\n".join(content_parts)
                    documents.append(Document(page_content=text, metadata={"source": file_path}))
            else:
                print(f"⚠️ Skipping unsupported file: {file_path}")
                continue
            print(f"  [✓] Loaded {file_path}")
        except Exception as e:
            print(f"❌ Failed to load {file_path}: {e}")
    return documents


def split_chunks(
    documents: list[Document],
    chunk_size: int = 800,
    chunk_overlap: int = 100,
    separators: list[str] | None = None,
    use_tokens: bool = True,
) -> list[Document]:
    """
    Split documents into semantic boundary chunks with token limits (512–800 tokens, 10–15% overlap)
    preserving markdown and table structures.
    Attaches page, chunk_index, source, and parent document_id to each chunk.
    """
    if separators is None:
        separators = [
            "\n## ",
            "\n### ",
            "\n#### ",
            "\n\n",
            "\n|",
            "\n",
            " ",
            "",
        ]
    if use_tokens:
        try:
            text_splitter = RecursiveCharacterTextSplitter.from_tiktoken_encoder(
                model_name="text-embedding-3-large",
                chunk_size=chunk_size,
                chunk_overlap=chunk_overlap,
                separators=separators,
            )
        except Exception:
            text_splitter = RecursiveCharacterTextSplitter(
                chunk_size=chunk_size * 4,
                chunk_overlap=chunk_overlap * 4,
                separators=separators,
            )
    else:
        text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
            separators=separators,
        )

    chunks = text_splitter.split_documents(documents)

    # Attach chunk metadata: page, chunk_index, source, document_id, chunk_id
    doc_chunk_counters: dict[str, int] = {}
    for chunk in chunks:
        src = str(chunk.metadata.get("source", "document"))
        p = Path(src)
        doc_id = str(chunk.metadata.get("document_id") or p.stem)
        idx = doc_chunk_counters.get(src, 0)
        doc_chunk_counters[src] = idx + 1

        chunk.metadata["document_id"] = doc_id
        chunk.metadata["source"] = src
        chunk.metadata["chunk_index"] = idx
        chunk.metadata["chunk_id"] = f"{doc_id}_chunk_{idx + 1}"
        if "page" not in chunk.metadata:
            chunk.metadata["page"] = chunk.metadata.get("page_number", 1)

    print(f"✂️ Created {len(chunks)} text chunks.")
    return chunks


def _is_retryable_embedding_error(exc: BaseException) -> bool:
    """Determine if an embedding API exception is transient (429 rate limit or 5xx server error)."""
    if isinstance(exc, (openai.RateLimitError, openai.APIConnectionError, openai.InternalServerError, openai.APITimeoutError)):
        return True
    status_code = getattr(exc, "status_code", None) or getattr(exc, "http_status", None)
    if status_code in (429, 500, 502, 503, 504):
        return True
    err_msg = str(exc).lower()
    return "429" in err_msg or "rate limit" in err_msg or "503" in err_msg or "timeout" in err_msg


@retry(
    retry=retry_if_exception(_is_retryable_embedding_error),
    wait=wait_exponential(multiplier=1, min=2, max=60),
    stop=stop_after_attempt(5),
    reraise=True,
)
def _embed_batch_with_retry(batch: list[str], model: OpenAIEmbeddings) -> list[list[float]]:
    return model.embed_documents(batch)


def embed_chunks_with_retry(
    texts: list[str],
    embeddings_model: OpenAIEmbeddings,
    batch_size: int = 64,
) -> list[list[float]]:
    """
    Batch chunk embeddings (64–128 items per call) with exponential backoff
    via tenacity for rate limits (429/503).
    """
    if not texts:
        return []
    embeddings: list[list[float]] = []
    total_batches = (len(texts) + batch_size - 1) // batch_size
    print(f"🚀 Generating embeddings for {len(texts)} chunks across {total_batches} batches (batch_size={batch_size})...")
    for i in range(0, len(texts), batch_size):
        batch = texts[i : i + batch_size]
        batch_embs = _embed_batch_with_retry(batch, embeddings_model)
        embeddings.extend(batch_embs)
    return embeddings


def get_embedding_model(api_key: str) -> OpenAIEmbeddings:
    return OpenAIEmbeddings(
        model="text-embedding-3-large",
        dimensions=1536,
        openai_api_key=api_key
    )


def rerank_candidates(
    query: str,
    candidates: list[dict],
    top_k: int = 5,
    cross_encoder_model: str | None = None,
) -> list[dict]:
    """
    Support optional cross-encoder re-ranking on top candidates.
    If a cross-encoder model is specified and available, scores candidates with it;
    otherwise applies robust deterministic fusion (combining lexical query overlap,
    semantic similarity, and reciprocal rank).
    """
    if not candidates:
        return []

    if cross_encoder_model:
        try:
            from sentence_transformers import CrossEncoder
            model = CrossEncoder(cross_encoder_model)
            pairs = [[query, c.get("content", "")] for c in candidates]
            scores = model.predict(pairs)
            for c, score in zip(candidates, scores):
                c["rerank_score"] = float(score)
            return sorted(candidates, key=lambda x: x["rerank_score"], reverse=True)[:top_k]
        except Exception as exc:
            logging.getLogger(__name__).warning("Cross-encoder model unavailable, applying fallback: %s", exc)

    # Deterministic hybrid scoring combining lexical token match, similarity, and RRF
    query_terms = set(re.findall(r"\w+", query.lower()))
    for c in candidates:
        content_terms = set(re.findall(r"\w+", c.get("content", "").lower()))
        term_overlap = len(query_terms & content_terms) / max(1, len(query_terms))
        base_sim = float(c.get("similarity", 0.0))
        rrf = float(c.get("rrf_score", 0.0))
        c["rerank_score"] = round(base_sim * 0.5 + term_overlap * 0.35 + rrf * 0.15, 4)

    return sorted(candidates, key=lambda x: x["rerank_score"], reverse=True)[:top_k]

def get_llm(api_key: str, model: str = "gpt-4o", temperature: float = 0.0) -> ChatOpenAI:
    return ChatOpenAI(model=model, temperature=temperature, openai_api_key=api_key)

def get_db_connection(database_url: str, timeout: int = 30):
    return psycopg2.connect(database_url, connect_timeout=timeout)

# --- Ollama utilities (extracted from py-scripts/ollama-models.py) ---------

def get_ollama_host(cli_host: str | None = None) -> str:
    """Return the Ollama host, preferring a CLI override, then the environment, then localhost."""
    host = cli_host or os.getenv("OLLAMA_HOST", "http://localhost:11434")
    return host.rstrip('/')


def get_timeout_seconds(default: int = 60) -> int:
    """Return the per-request timeout from the environment or a safer default."""
    raw_value = os.getenv("OLLAMA_TIMEOUT_SECONDS", str(default))
    try:
        timeout = int(raw_value)
        if timeout <= 0:
            raise ValueError("timeout must be positive")
        return timeout
    except (TypeError, ValueError):
        print(f"[WARN] Invalid OLLAMA_TIMEOUT_SECONDS='{raw_value}'. Using default {default} seconds.")
        return default


def get_retry_attempts(default: int = 1) -> int:
    """Return how many times a timed-out request should be retried."""
    raw_value = os.getenv("OLLAMA_RETRY_ATTEMPTS", str(default))
    try:
        attempts = int(raw_value)
        if attempts <= 0:
            raise ValueError("retry attempts must be positive")
        return attempts
    except (TypeError, ValueError):
        print(f"[WARN] Invalid OLLAMA_RETRY_ATTEMPTS='{raw_value}'. Using default {default} attempt.")
        return default

# --- Logging utility (extracted from py-pgvector-local/api.py) -------------

def log_error(stage: str, error: Exception) -> None:
    """Log an error message with stage context and traceback."""
    import traceback
    print(f"[ERROR] {stage}: {error}")
    traceback.print_exc()

def safe_calculate(expression: str) -> str:
    """Evaluate a mathematical expression safely using AST parsing.
    Supports: +, -, *, /, ** and parentheses.
    """
    try:
        tree = ast.parse(expression.strip(), mode="eval")
        def _eval(node):
            match node:
                case ast.Constant(value=v) if isinstance(v, int | float):
                    return v
                case ast.BinOp(left=left, op=op, right=right) if type(op) in _SAFE_OPERATORS:
                    return _SAFE_OPERATORS[type(op)](_eval(left), _eval(right))
                case ast.UnaryOp(op=op, operand=operand) if type(op) in _SAFE_OPERATORS:
                    return _SAFE_OPERATORS[type(op)](_eval(operand))
                case _:
                    raise ValueError(f"Unsupported expression: {ast.dump(node)}")
        result = _eval(tree.body)
        return str(result)
    except (ValueError, ZeroDivisionError, SyntaxError) as e:
        return f"Error evaluating expression: {e}"

def remove_comments_and_docstrings(filepath: str) -> None:
    """Remove comments and docstrings from a source code file."""
    if not os.path.exists(filepath):
        print(f"File not found: {filepath}")
        return
    with open(filepath, 'r', encoding='utf-8') as file:
        content = file.read()
    content = _DOCSTRING_PATTERN.sub('', content)
    content = re.sub(r'--.*', '', content)
    content = re.sub(r'#.*', '', content)
    lines = [line for line in content.split('\n') if line.strip()]
    with open(filepath, 'w', encoding='utf-8') as file:
        file.write('\n'.join(lines))
    print(f"Cleaned: {filepath}")
