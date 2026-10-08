import tempfile
from pathlib import Path
from unittest.mock import MagicMock, patch

import docx
import pytest
from langchain_core.documents import Document

from libs.utilities import get_documents


def test_get_documents_with_docx(tmp_path):
    doc_path = tmp_path / "sample.docx"
    doc = docx.Document()
    doc.add_heading("Architecture Overview", level=1)
    doc.add_paragraph("This document outlines system architecture.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Service"
    table.cell(0, 1).text = "Port"
    table.cell(1, 0).text = "API Gateway"
    table.cell(1, 1).text = "8000"
    doc.save(str(doc_path))

    docs = get_documents([str(doc_path)])
    assert len(docs) == 1
    content = docs[0].page_content
    assert "Architecture Overview" in content
    assert "This document outlines system architecture." in content
    assert "Service | Port" in content
    assert "API Gateway | 8000" in content
    assert docs[0].metadata.get("source") == str(doc_path)


def test_get_documents_docx_with_docx2txt_loader(tmp_path):
    doc_path = tmp_path / "mocked.docx"
    doc_path.write_text("dummy")

    mock_doc = Document(page_content="Mocked docx2txt content", metadata={"source": str(doc_path)})
    with patch("libs.utilities.Docx2txtLoader") as mock_loader_cls:
        instance = MagicMock()
        instance.load.return_value = [mock_doc]
        mock_loader_cls.return_value = instance

        docs = get_documents([str(doc_path)])
        assert len(docs) == 1
        assert docs[0].page_content == "Mocked docx2txt content"
        mock_loader_cls.assert_called_once_with(str(doc_path))


def test_get_documents_skips_unsupported(tmp_path):
    txt_path = tmp_path / "sample.unsupported"
    txt_path.write_text("Hello")

    docs = get_documents([str(txt_path)])
    assert len(docs) == 0


def test_get_documents_handles_invalid_file(tmp_path):
    broken_docx = tmp_path / "broken.docx"
    broken_docx.write_text("not a valid docx zip content")

    docs = get_documents([str(broken_docx)])
    assert len(docs) == 0


def test_compute_file_hash_deterministic(tmp_path):
    import hashlib
    from libs.utilities import compute_file_hash

    test_file = tmp_path / "test_doc.md"
    content = b"# Architecture Blueprint\nContent for cryptographic de-duplication testing."
    test_file.write_bytes(content)

    expected_hash = hashlib.sha256(content).hexdigest()
    computed_hash = compute_file_hash(test_file)

    assert computed_hash == expected_hash
    assert len(computed_hash) == 64


def test_split_chunks_attaches_semantic_metadata(tmp_path):
    from libs.utilities import split_chunks

    doc = Document(
        page_content=(
            "# Header 1\nSection text about RBAC.\n\n"
            "## Table Section\n"
            "| Column 1 | Column 2 |\n"
            "| Value A  | Value B  |\n\n"
            "Additional details about security architecture."
        ),
        metadata={"source": "docs/architecture.md", "page": 2}
    )

    chunks = split_chunks([doc], chunk_size=50, chunk_overlap=10)
    assert len(chunks) >= 1
    for idx, chunk in enumerate(chunks):
        assert chunk.metadata["source"] == "docs/architecture.md"
        assert chunk.metadata["document_id"] == "architecture"
        assert chunk.metadata["chunk_index"] == idx
        assert chunk.metadata["chunk_id"] == f"architecture_chunk_{idx + 1}"
        assert chunk.metadata["page"] == 2


def test_embed_chunks_with_retry():
    import openai
    from libs.utilities import embed_chunks_with_retry

    mock_model = MagicMock()
    call_count = 0

    def mock_embed(batch):
        nonlocal call_count
        call_count += 1
        if call_count == 1:
            # Simulate transient 429 rate limit error on first attempt
            err = openai.RateLimitError(
                message="Rate limit exceeded",
                response=MagicMock(status_code=429, headers={}),
                body=None
            )
            raise err
        return [[0.1] * 1536 for _ in batch]

    mock_model.embed_documents.side_effect = mock_embed

    texts = ["Text chunk 1", "Text chunk 2", "Text chunk 3"]
    # With exponential backoff, it should retry and succeed on attempt 2
    embeddings = embed_chunks_with_retry(texts, mock_model, batch_size=2)
    assert len(embeddings) == 3
    assert all(len(emb) == 1536 for emb in embeddings)
    assert call_count >= 2
